"""Generate short impact workbooks: problem, what we have, two-way fix.

Scope: only columns V4 ingested from the old schema (legacy_ehr_v1).
Fixes are written two ways: claims-table-only (first run) vs all tables.
"""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT_DIR = r"c:\Users\PAMALI\Desktop\RDDT\results\analysis"

NAVY = PatternFill("solid", fgColor="1F4E79")
SECTION = PatternFill("solid", fgColor="2F5597")
SUB = PatternFill("solid", fgColor="D6E3F0")
RED = PatternFill("solid", fgColor="FCE4D6")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="E2EFDA")
WHITE = PatternFill("solid", fgColor="FFFFFF")

WHITE_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
BODY = Font(name="Calibri", size=10)
BOLD = Font(name="Calibri", size=10, bold=True)
RED_FONT = Font(name="Calibri", size=10, bold=True, color="C00000")
ORANGE_FONT = Font(name="Calibri", size=10, bold=True, color="B25900")
GREEN_FONT = Font(name="Calibri", size=10, bold=True, color="375623")
ITALIC = Font(name="Calibri", size=9, italic=True, color="595959")

THIN = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
WRAP = Alignment(wrap_text=True, vertical="center")
WRAP_LEFT = Alignment(wrap_text=True, vertical="center", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

STATUS_FILL = {
    "BLOCKER": RED, "HIGH": RED, "MISSING": RED, "UNMAPPED": RED,
    "100% NULL": RED, "100% EMPTY": RED, "TABLE MISSING": RED, "CANNOT FIX": RED,
    "UNVERIFIED": YELLOW, "PARTIAL": YELLOW, "RENAMED": YELLOW, "WRONG FORMAT": YELLOW,
    "WRONG CODES": YELLOW, "LOW": GREEN, "OK": GREEN, "POPULATED": GREEN,
    "MAPPED": GREEN, "YES": GREEN, "NO": RED, "N/A": WHITE, "EMPTY": RED,
    "IGNORE": WHITE, "NOTE": YELLOW,
}
STATUS_FONT = {
    "BLOCKER": RED_FONT, "HIGH": RED_FONT, "MISSING": RED_FONT, "UNMAPPED": RED_FONT,
    "100% NULL": RED_FONT, "100% EMPTY": RED_FONT, "TABLE MISSING": RED_FONT,
    "CANNOT FIX": RED_FONT, "NO": RED_FONT,
    "UNVERIFIED": ORANGE_FONT, "PARTIAL": ORANGE_FONT, "RENAMED": ORANGE_FONT,
    "WRONG FORMAT": ORANGE_FONT, "WRONG CODES": ORANGE_FONT, "NOTE": ORANGE_FONT,
    "OK": GREEN_FONT, "POPULATED": GREEN_FONT, "MAPPED": GREEN_FONT, "YES": GREEN_FONT,
}

SCOPE = (
    "Left side = fields the algorithm expected. "
    "Right side = whether that content is in the current data. "
    "No guessed labels for empty CLAIMS slots. "
    "Claims-only = first run on CLAIMS alone. All-tables = every in-scope table we have."
)


def _style(cell, header=False, fill=None):
    cell.border = THIN
    cell.alignment = CENTER if header else WRAP
    if header:
        cell.fill = NAVY
        cell.font = HEADER_FONT
        return
    val = str(cell.value or "").strip()
    cell.fill = STATUS_FILL.get(val, fill or WHITE)
    cell.font = STATUS_FONT.get(val, BODY)


def _widths(ws: Worksheet, widths: list[int]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws: Worksheet, text: str, cols: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    cell = ws.cell(1, 1, text)
    cell.fill = SECTION
    cell.font = TITLE_FONT
    cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 28


def _note(ws: Worksheet, row: int, text: str, cols: int, height=32):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    cell.font = ITALIC
    cell.alignment = WRAP_LEFT
    cell.fill = SUB
    ws.row_dimensions[row].height = height


def _headers(ws: Worksheet, row: int, headers: list[str]):
    for i, h in enumerate(headers, 1):
        _style(ws.cell(row, i, h), header=True)
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = f"A{row + 1}"
    return row + 1


def _row(ws: Worksheet, row: int, values: list, height=44):
    for i, v in enumerate(values, 1):
        _style(ws.cell(row, i, v))
    ws.row_dimensions[row].height = height
    return row + 1


def add_sheet(wb: Workbook, title: str, heading: str, note: str, headers: list[str],
              rows: list[list], widths: list[int], height=48) -> Worksheet:
    ws = wb.active if wb.active.title == "Sheet" and len(wb.sheetnames) == 1 else wb.create_sheet()
    ws.title = title
    _title(ws, heading, len(headers))
    _note(ws, 2, note, len(headers))
    r = _headers(ws, 3, headers)
    for row in rows:
        r = _row(ws, r, row, height)
    _widths(ws, widths)
    ws.sheet_view.showGridLines = False
    return ws


# table, expected field, old column, in_data (YES/NO), what the data shows
EXPECTED = [
    ["CLAIMS", "patient_id", "Member/PatientId", "YES",
     "Present. Patient GUIDs (COLUMN0). 929,356 patients."],
    ["CLAIMS", "encounter_id", "EncounterId/VisitId", "YES",
     "Present. Numeric visit-like ids (COLUMN1). Join to ENCOUNTERS.VISITID not counted."],
    ["CLAIMS", "diagnosis_type", "DiagnosisType", "YES",
     "Present. Every row is ICD-10-CM (COLUMN7)."],
    ["CLAIMS", "diagnosis_code", "DiagnosisCode", "YES",
     "Present. ICD-10 codes (COLUMN8). 134 E85 patients already found."],
    ["CLAIMS", "other_diagnosis_9", "OtherDiagnosisCodes9", "NO",
     "Not in the data. No extra ICD-9 values on the claim rows."],
    ["CLAIMS", "other_diagnosis_10", "OtherDiagnosisCodes10", "NO",
     "Not in the data. No extra ICD-10 values on the claim rows."],
    ["CLAIMS", "procedure_code", "ProcedureCode", "NO",
     "Not in the data. No CPT/HCPCS values on the claim rows."],
    ["CLAIMS", "clinical_notes", "ClinicalNotes", "NO",
     "Not in the data. No note text on the claim rows."],
    ["CLAIMS", "from_date", "FromDate", "NO",
     "Not in the data. No date values on the claim rows."],
    ["CLAIMS", "to_date", "ToDate", "NO",
     "Not in the data. No date values on the claim rows."],
    ["ENCOUNTERS", "encounter_id", "EncounterId/VisitId", "YES",
     "Present. VISITID. 16,334,743 distinct. Some duplicate rows."],
    ["ENCOUNTERS", "patient_id", "Member/PatientId", "YES",
     "Present. PATIENTID. 1,056,034 patients."],
    ["ENCOUNTERS", "encounter_date", "Encounter/Visit Date", "YES",
     "Present. VISITDATE. Real dates. 1 null. 38,878 future dates."],
    ["CENSUS", "patient_id", "Member/PatientId", "YES",
     "Present. PATIENTID. 2,794,815 patients."],
    ["CENSUS", "birth_date", "BirthDate", "YES",
     "Present. BIRTHDATE. 3.36% missing. 96% of the rest are age 51-60."],
    ["CENSUS", "gender", "Gender", "YES",
     "Present. GENDER. F/M/U filled."],
    ["CENSUS", "family_id", "FamilyId", "NO",
     "Not in the data. No family id column on CENSUS."],
    ["LABS", "patient_id", "Member/PatientId", "YES",
     "Present. PATIENTID. 695,413 patients."],
    ["LABS", "encounter_id", "EncounterId/VisitId", "YES",
     "Present. ENCOUNTERKEY."],
    ["LABS", "lab_id", "LabId", "YES",
     "Present. LABTESTKEY."],
    ["LABS", "lab_result_id", "LabResultId", "NO",
     "Not in the data. LABRESULTID is empty."],
    ["LABS", "observation_identifier", "ObservationIdentifier", "YES",
     "Present as test names (Troponin I, High Sensitivity; BNP), not LOINC codes. OBSERVATIONIDENTIFIER."],
    ["LABS", "observation_value", "ObservationValue", "YES",
     "Present. OBSERVATIONVALUE. Numeric/text results."],
    ["LABS", "result_status", "ObservationResultStatus", "YES",
     "Present. Mostly Final result."],
    ["LABS", "observation_datetime", "ObservationDateTime", "YES",
     "Present. OBSERVATIONDATETIME. Real timestamps. 0% null."],
    ["LABS", "lab_result_note", "LabResultNote", "YES",
     "Present but operational comments, not ATTR narrative. LABRESULTNOTE."],
    ["SURGICAL_HISTORY", "patient_id", "Member/PatientId", "YES",
     "Present. PATIENTID. 8,329 patients only."],
    ["SURGICAL_HISTORY", "encounter_id", "EncounterId/VisitId", "YES",
     "Present. ENCOUNTERID. Mostly the word Unknown."],
    ["SURGICAL_HISTORY", "record_id", "SurgicalHistoryId", "YES",
     "Present. SURGICALHISTORYID."],
    ["SURGICAL_HISTORY", "value", "Value", "YES",
     "Present. CATEGORY. 92% is the generic label SURGICAL HISTORY PROCEDURES."],
    ["SURGICAL_HISTORY", "snomed", "SNOMED", "YES",
     "Present. SNOMED filled. 48 codes. None are in the ATTR catalog."],
    ["SURGICAL_HISTORY", "secondary_snomed", "Secondary SNOMED", "NO",
     "Not in the data. SECONDARY_SNOMED is empty."],
    ["SURGICAL_HISTORY", "event_date", "Date", "YES",
     "Present. DATE. 7.53% null. Rest 2020-2026."],
    ["CLINICAL_NOTE", "note_text, note_id, dates", "Clinical Note Text, NoteId, Date", "NO",
     "Not in the data. Table is missing."],
    ["MEDICAL_HISTORY", "snomed, value, date", "SNOMED, Value, Date", "NO",
     "Not in the data. Table is missing."],
    ["FAMILY_HISTORY", "snomed, condition, family_member, date", "SNOMED, Condition, FamilyMember, Date", "NO",
     "Not in the data. Table is missing."],
]

PROBLEM_HAVE_DO = [
    ["Claim dates are not in the data, so suspicion cannot score.",
     "Patient id, encounter id, ICD-10 type, and ICD-10 code are present. Dates, CPT, notes, and extra diagnoses are not.",
     "Run confirmed-only (already works). For suspicion: score ICD-10 without recency. Cannot timestamp from claims alone.",
     "If claim encounter ids join to ENCOUNTERS.VISITID, use VISITDATE. Labs and surgical already have dates."],
    ["Procedure codes are not in the data.",
     "No CPT/HCPCS values on claims.",
     "Use ICD-10 diagnosis atoms only. Do not treat a diagnosis as a surgery.",
     "Still no CPT. Surgical SNOMED is present but not ATTR procedures."],
    ["Notes are not in the data.",
     "No note text on claims. CLINICAL_NOTE table is missing.",
     "No NLP. ICD-10 only.",
     "Still no notes. Do not use lab comments as a substitute."],
    ["Extra diagnoses are not in the data.",
     "Only the primary ICD-10 code is present.",
     "Score that primary code. Other claim lines for the same patient still count if they are primary on those lines.",
     "Same. MEDICAL_HISTORY is missing, so no extra conditions from history."],
    ["Lab identifiers are names, not LOINC codes.",
     "Test names, values, and dates are present (Troponin I HS, BNP, NT-proBNP).",
     "N/A. Claims-only does not use labs.",
     "Match the names (and values for elevated). Leave labs off until that mapping exists."],
    ["Surgical SNOMED is present but not ATTR.",
     "48 SNOMED codes (cesarean, transplant, orderables). None in the ATTR catalog. 8,329 patients.",
     "N/A.",
     "Leave surgical_history off. It does not replace missing CPT."],
    ["Notes / medical history / family history tables are missing.",
     "Those three tables are not in the warehouse. Census family id is also not in the data.",
     "Expected for a claims-first run. Those routes stay off.",
     "Those routes cannot proceed until the tables return."],
    ["Census family id is not in the data. Age is not useful.",
     "Patient id, gender, and birth date are present. 96% of ages sit in 51-60.",
     "N/A.",
     "Do not rank by age. No pedigree from census."],
]

SHORT_PROBLEMS = [
    ["CLAIMS dates", "BLOCKER", "No date values on claims",
     "Confirmed E85 works. Suspicion does not.",
     "Presence-only ICD-10, or confirmed-only.",
     "Join claim encounter id to VISITID; use VISITDATE."],
    ["CLAIMS procedures", "HIGH", "No CPT values",
     "No procedure atoms.",
     "ICD-10 diagnoses only.",
     "Still no CPT. Surgical SNOMED does not help."],
    ["CLAIMS notes", "HIGH", "No note text",
     "No claim NLP.",
     "ICD-10 only.",
     "CLINICAL_NOTE table missing. No NLP."],
    ["CLAIMS extra dx", "LOW", "No extra diagnosis values",
     "Only primary ICD-10.",
     "Score the primary code.",
     "MEDICAL_HISTORY missing."],
    ["ENCOUNTERS dates", "UNVERIFIED", "VISITDATE is present; join not counted",
     "Only dated table that could stamp claims.",
     "Do not use in a claims-only run.",
     "Count the join, then backfill claim dates."],
    ["LABS identifiers", "HIGH", "Names, not LOINC",
     "Engine looks for codes; data has names.",
     "Do not use labs.",
     "Map names + values. Do not enable until then."],
    ["SURGICAL_HISTORY SNOMED", "LOW", "Wrong codes for ATTR",
     "Almost no ATTR signal.",
     "Do not use.",
     "Leave disabled."],
    ["CLINICAL_NOTE", "BLOCKER", "Table missing",
     "No note NLP.",
     "Expected. Skip.",
     "Cannot proceed with notes until the table exists."],
    ["MEDICAL_HISTORY", "BLOCKER", "Table missing",
     "No history SNOMED.",
     "Expected. Skip.",
     "Cannot proceed with history until the table exists."],
    ["FAMILY_HISTORY", "BLOCKER", "Table missing",
     "No family atoms.",
     "Expected. Skip.",
     "Cannot proceed with family until the table exists."],
]


# Incoming warehouse columns the algorithm does not read.
# table, column, in_data, used_by_algorithm, what to do
EXTRA = [
    ["CENSUS", "DEATHDATE", "EMPTY", "NO",
     "Column exists. First 10 Snowflake rows and the 50-row sample are blank. Prior census column review: 100% blank strings. The census notebook did not run a dedicated COUNT on DEATHDATE. If it were filled, drop deceased patients from screening. Today we cannot exclude anyone."],
    ["CENSUS", "OMBETHNICITY / OMBRACE", "YES", "NO",
     "Filled in the sample (e.g. Hispanic or Latino, White). Not used by the algorithm. Optional profile only."],
    ["CENSUS", "FIRSTNAME, LASTNAME, ADDRESS, CITY, STATE, ZIP, etc.", "EMPTY", "NO",
     "Present as columns, blank in the sample. Ignore."],
    ["LABS", "OBSERVATIONIDENTIFIERCOMMONNAME", "YES", "NO",
     "Short alias of OBSERVATIONIDENTIFIER (already mapped). Same test, shorter label. Do not remap; it does not add a new clinical fact."],
    ["LABS", "PERFORMEDDATETIME", "YES", "NO",
     "Second timestamp. OBSERVATIONDATETIME is already mapped and populated. Ignore."],
    ["LABS", "ANALYSISDATETIME", "PARTIAL", "NO",
     "Third timestamp. Not needed for matching."],
    ["ENCOUNTERS", "TYPE", "YES", "NO",
     "Visit label (Inpatient, FOLLOW UP, echo, LAB). Not an atom. Ignore unless we later want setting context."],
    ["ENCOUNTERS", "FACILITYID, PHYSICIANID, FINANCIALCLASS, STATUS", "YES", "NO",
     "Hospital / provider / insurance. Out of scope."],
    ["SURGICAL_HISTORY", "SOURCE", "YES", "NO",
     "Metadata (Surgical History Procedure). Not an atom."],
    ["CLAIMS", "COLUMN2, COLUMN3, COLUMN4, COLUMN9", "YES", "NO",
     "Claim number, line, Closed status, NPI-like values. Admin / provider. Out of scope."],
    ["MEDICATIONS", "(entire table)", "YES", "NO",
     "Out of scope per request. Not analyzed."],
]


def _filter_map(*tables: str) -> list[list]:
    return [row for row in EXPECTED if row[0] in tables]


def _filter_extra(*tables: str) -> list[list]:
    return [row for row in EXTRA if row[0] in tables]


HAVE_DO_HEADERS = [
    "The problem",
    "What we have",
    "Fix - claims table only",
    "Fix - all tables",
]
HAVE_DO_WIDTHS = [32, 44, 44, 44]

MAP_HEADERS = [
    "Table",
    "Expected field",
    "Old column",
    "In the data?",
    "What the data shows",
]
MAP_WIDTHS = [22, 28, 28, 14, 72]
MAP_NOTE = (
    "Expected = what the algorithm needs. "
    "YES/NO is from the values, not from guessed empty-column names. "
    "CLAIMS COLUMN0/1/7/8 are named only because the values identify them."
)
EXTRA_HEADERS = [
    "Table",
    "Incoming column (not in the algorithm)",
    "In the data?",
    "Used by algorithm?",
    "Note",
]
EXTRA_WIDTHS = [20, 36, 14, 18, 72]
EXTRA_NOTE = (
    "These columns exist in the incoming warehouse and the algorithm does not read them. "
    "Medications are out of scope. Empty CLAIMS slots are not listed."
)


def build_master(wb: Workbook):
    add_sheet(wb, "00_Problem_Have_Do",
              "The problem, what we have, what we can do",
              SCOPE, HAVE_DO_HEADERS, PROBLEM_HAVE_DO, HAVE_DO_WIDTHS, 62)
    add_sheet(wb, "01_Expected_vs_Data",
              "Expected fields vs what is in the data",
              MAP_NOTE, MAP_HEADERS, EXPECTED, MAP_WIDTHS, 32)
    add_sheet(wb, "02_Problems_and_Fixes",
              "Each problem: cause, claims-only fix, all-tables fix",
              "First run is claims-only. All-tables uses every present in-scope table.",
              ["Problem", "Severity", "Caused by", "What breaks", "Fix - claims only", "Fix - all tables"],
              SHORT_PROBLEMS, [22, 14, 28, 32, 36, 40], 40)
    add_sheet(wb, "03_Extra_Incoming_Fields",
              "Incoming fields the algorithm does not use",
              EXTRA_NOTE, EXTRA_HEADERS, EXTRA, EXTRA_WIDTHS, 48)
    return wb


def build_claims(wb: Workbook, active=True):
    add_sheet(wb, "01_Problem_Have_Do",
              "CLAIMS - problem, what we have, what we can do",
              "First-run table. 53,311,842 lines / 929,356 patients.",
              HAVE_DO_HEADERS, PROBLEM_HAVE_DO[0:4], HAVE_DO_WIDTHS, 58)
    add_sheet(wb, "02_Expected_vs_Data",
              "CLAIMS - expected vs data",
              MAP_NOTE, MAP_HEADERS, _filter_map("CLAIMS"), MAP_WIDTHS, 34)
    add_sheet(wb, "03_Problems_and_Fixes",
              "CLAIMS - fixes two ways",
              "Claims-only cannot borrow encounter dates. All-tables can, if the visit join holds.",
              ["Problem", "Severity", "Caused by", "What breaks", "Fix - claims only", "Fix - all tables"],
              [row for row in SHORT_PROBLEMS if row[0].startswith("CLAIMS")],
              [22, 14, 28, 32, 36, 40], 44)
    add_sheet(wb, "04_Extra_Incoming_Fields",
              "CLAIMS - extra incoming fields",
              EXTRA_NOTE, EXTRA_HEADERS, _filter_extra("CLAIMS"), EXTRA_WIDTHS, 40)
    return wb


def build_encounters(wb: Workbook, active=True):
    add_sheet(wb, "01_Problem_Have_Do",
              "ENCOUNTERS - problem, what we have, what we can do",
              "Date lookup only. Not an atom source.",
              HAVE_DO_HEADERS,
              [[
                  "Claims have no dates; encounters do.",
                  "VISITID, PATIENTID, VISITDATE are all present.",
                  "Do not use this table in a claims-only run.",
                  "If claim encounter ids match VISITID, use VISITDATE on claims. Drop future dates. Deduplicate VISITID.",
              ]],
              HAVE_DO_WIDTHS, 70)
    add_sheet(wb, "02_Expected_vs_Data",
              "ENCOUNTERS - expected vs data",
              "Facility, physician, visit type were never algorithm fields and are omitted.",
              MAP_HEADERS, _filter_map("ENCOUNTERS"), MAP_WIDTHS, 40)
    add_sheet(wb, "03_Problems_and_Fixes",
              "ENCOUNTERS - fixes two ways",
              SCOPE,
              ["Problem", "Severity", "Caused by", "What breaks", "Fix - claims only", "Fix - all tables"],
              [row for row in SHORT_PROBLEMS if row[0].startswith("ENCOUNTERS")],
              [22, 14, 28, 32, 36, 40], 44)
    add_sheet(wb, "04_Extra_Incoming_Fields",
              "ENCOUNTERS - extra incoming fields",
              EXTRA_NOTE, EXTRA_HEADERS, _filter_extra("ENCOUNTERS"), EXTRA_WIDTHS, 40)
    return wb


def build_census(wb: Workbook, active=True):
    add_sheet(wb, "01_Problem_Have_Do",
              "CENSUS - problem, what we have, what we can do",
              "Algorithm fields only: patient_id, birth_date, gender, family_id.",
              HAVE_DO_HEADERS, PROBLEM_HAVE_DO[7:8], HAVE_DO_WIDTHS, 70)
    add_sheet(wb, "02_Expected_vs_Data",
              "CENSUS - expected vs data",
              "Name, address, race, city, state exist in the warehouse and are omitted.",
              MAP_HEADERS, _filter_map("CENSUS"), MAP_WIDTHS, 36)
    add_sheet(wb, "03_Problems_and_Fixes",
              "CENSUS - fixes two ways",
              "Census is not used in the claims-only first run.",
              ["Problem", "Severity", "Caused by", "What breaks", "Fix - claims only", "Fix - all tables"],
              [["Census demographics", "LOW", "Family id not in the data; age clustered 51-60",
                "No pedigree. Age not a useful gate.", "N/A",
                "Optional overlay. Missing DOB = UNKNOWN. Do not rank by age."],
               ["Death date", "NOTE", "DEATHDATE column exists but is blank in the analysis files",
                "Cannot tell who is deceased, so we cannot drop dead patients.",
                "N/A today.",
                "If DEATHDATE is ever filled, exclude those patients from screening. Re-count before using as a filter."]],
              [22, 14, 28, 32, 36, 40], 48)
    add_sheet(wb, "04_Extra_Incoming_Fields",
              "CENSUS - extra incoming fields",
              "DEATHDATE is the only extra field that would change screening if it filled.",
              EXTRA_HEADERS, _filter_extra("CENSUS"), EXTRA_WIDTHS, 52)
    return wb


def build_labs(wb: Workbook, active=True):
    add_sheet(wb, "01_Problem_Have_Do",
              "LABS - problem, what we have, what we can do",
              "133.9M rows / 695k patients.",
              HAVE_DO_HEADERS, PROBLEM_HAVE_DO[4:5], HAVE_DO_WIDTHS, 70)
    add_sheet(wb, "02_Expected_vs_Data",
              "LABS - expected vs data",
              "PERFORMEDDATETIME, ANALYSISDATETIME, COMMONNAME, LABREQUESTID omitted.",
              MAP_HEADERS, _filter_map("LABS"), MAP_WIDTHS, 36)
    add_sheet(wb, "03_Problems_and_Fixes",
              "LABS - fixes two ways",
              "BNP/troponin atoms have no LOINC entries in the catalog.",
              ["Problem", "Severity", "Caused by", "What breaks", "Fix - claims only", "Fix - all tables"],
              [row for row in SHORT_PROBLEMS if row[0].startswith("LABS")],
              [22, 14, 28, 32, 36, 40], 48)
    add_sheet(wb, "04_Extra_Incoming_Fields",
              "LABS - extra incoming fields",
              "COMMONNAME is a short alias of the test name already mapped as observation_identifier.",
              EXTRA_HEADERS, _filter_extra("LABS"), EXTRA_WIDTHS, 48)
    return wb


def build_surgical(wb: Workbook, active=True):
    add_sheet(wb, "01_Problem_Have_Do",
              "SURGICAL_HISTORY - problem, what we have, what we can do",
              "SNOMED, text, date, ids. SOURCE omitted.",
              HAVE_DO_HEADERS, PROBLEM_HAVE_DO[5:6], HAVE_DO_WIDTHS, 70)
    add_sheet(wb, "02_Expected_vs_Data",
              "SURGICAL_HISTORY - expected vs data",
              MAP_NOTE, MAP_HEADERS, _filter_map("SURGICAL_HISTORY"), MAP_WIDTHS, 36)
    add_sheet(wb, "03_Problems_and_Fixes",
              "SURGICAL_HISTORY - fixes two ways",
              SCOPE,
              ["Problem", "Severity", "Caused by", "What breaks", "Fix - claims only", "Fix - all tables"],
              [row for row in SHORT_PROBLEMS if row[0].startswith("SURGICAL")],
              [22, 14, 28, 32, 36, 40], 48)
    add_sheet(wb, "04_Extra_Incoming_Fields",
              "SURGICAL_HISTORY - extra incoming fields",
              EXTRA_NOTE, EXTRA_HEADERS, _filter_extra("SURGICAL_HISTORY"), EXTRA_WIDTHS, 40)
    return wb


def build_missing(wb: Workbook, active=True):
    add_sheet(wb, "01_Problem_Have_Do",
              "Missing tables - cannot proceed on these routes",
              "Social history and medications are out of scope.",
              HAVE_DO_HEADERS, PROBLEM_HAVE_DO[6:7], HAVE_DO_WIDTHS, 62)
    add_sheet(wb, "02_Expected_vs_Data",
              "Missing tables - expected vs data",
              "The table is gone, so every expected field is not in the data.",
              MAP_HEADERS,
              _filter_map("CLINICAL_NOTE", "MEDICAL_HISTORY", "FAMILY_HISTORY"),
              MAP_WIDTHS, 40)
    add_sheet(wb, "03_Problems_and_Fixes",
              "Missing tables - fixes two ways",
              "Claims-first run never needed these.",
              ["Problem", "Severity", "Caused by", "What breaks", "Fix - claims only", "Fix - all tables"],
              [row for row in SHORT_PROBLEMS if row[0] in {"CLINICAL_NOTE", "MEDICAL_HISTORY", "FAMILY_HISTORY"}],
              [22, 14, 28, 32, 36, 40], 40)
    return wb


def save_book(builder, filename: str):
    wb = Workbook()
    builder(wb, active=True)
    path = os.path.join(OUT_DIR, filename)
    wb.save(path)
    print("saved", path)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    master = Workbook()
    build_master(master)
    path = os.path.join(OUT_DIR, "schema_and_data_impact_analysis.xlsx")
    master.save(path)
    print("saved", path)

    save_book(build_claims, "claims_schema_and_data_analysis.xlsx")
    save_book(build_encounters, "encounters_schema_and_data_analysis.xlsx")
    save_book(build_census, "census_schema_and_data_analysis.xlsx")
    save_book(build_labs, "labs_schema_and_data_analysis.xlsx")
    save_book(build_surgical, "surgical_history_schema_and_data_analysis.xlsx")
    save_book(build_missing, "missing_tables_impact_analysis.xlsx")


if __name__ == "__main__":
    main()
