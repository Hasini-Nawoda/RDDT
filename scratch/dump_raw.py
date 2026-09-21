import openpyxl, glob, os, sys

for path in sorted(glob.glob(r"c:\Users\PAMALI\Desktop\RDDT\results\raw\*.xlsx")):
    print("=" * 100)
    print("FILE:", os.path.basename(path))
    print("=" * 100)
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print(f"\n--- SHEET: {ws.title}  ({ws.max_row} rows x {ws.max_column} cols) ---")
        n = 0
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            while vals and vals[-1] == "":
                vals.pop()
            if not vals:
                continue
            line = " | ".join(v[:150] for v in vals)
            print(line)
            n += 1
            if n > 120:
                print(f"... ({ws.max_row - n} more rows truncated)")
                break
    wb.close()
