"""Dump Excel workbook contents as text for analysis."""
import os
from openpyxl import load_workbook

OUT = r"c:\Users\PAMALI\Desktop\RDDT\scratch\excel_dumps"
os.makedirs(OUT, exist_ok=True)

files = [
    r"c:\Users\PAMALI\Desktop\RDDT\results\raw\schema_counts_summary_outputs.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\raw\census_data_analysis_outputs.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\raw\claim_data_analysis_outputs.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\raw\encounter_visit_data_analysis_outputs.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\raw\lab_data_analysis_outputs.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\raw\surgical_history_data_analysis_outputs.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\raw\clinical_note_data_analysis_outputs.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\raw\family_history_data_analysis_outputs.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\raw\medical_history_data_analysis_outputs.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\analysis\schema_and_data_impact_analysis.xlsx",
    r"c:\Users\PAMALI\Desktop\RDDT\results\analysis\missing_tables_impact_analysis.xlsx",
]


def dump_sheet(ws, max_rows=80, max_cols=20):
    rows = []
    for i, row in enumerate(ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True), 1):
        if all(v is None for v in row):
            continue
        cells = []
        for v in row:
            if v is None:
                cells.append("")
            else:
                s = str(v).replace("\n", " | ")
                if len(s) > 400:
                    s = s[:400] + "..."
                cells.append(s)
        rows.append(f"R{i}: " + " || ".join(cells))
    return "\n".join(rows)


for path in files:
    wb = load_workbook(path, read_only=True, data_only=True)
    name = os.path.splitext(os.path.basename(path))[0]
    out_path = os.path.join(OUT, name + ".txt")
    parts = [f"FILE: {path}", f"SHEETS: {wb.sheetnames}", ""]
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append("=" * 80)
        parts.append(f"SHEET: {sheet_name}")
        parts.append("=" * 80)
        max_rows = 200 if sheet_name.lower() in {"run summary", "run_summary", "executive_summary"} else 80
        parts.append(dump_sheet(ws, max_rows=max_rows, max_cols=16))
        parts.append("")
    wb.close()
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))
    print("wrote", out_path)
