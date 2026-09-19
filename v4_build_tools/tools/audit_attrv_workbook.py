from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl


REQUIRED_SHEETS = {
    "Signals",
    "Signal_Atoms",
    "Atoms",
    "Terminology",
    "Buckets",
    "Combinations",
    "Combination_Requirements",
    "Requirement_Buckets",
    "Requirement_Tiers",
    "Requirement_Signals",
    "Priority_Policies",
    "Combination_Rules",
    "Combination_Rule_Groups",
    "Combination_Rule_Members",
    "Guardrails",
    "Algorithm_Contract",
    "Controlled_Vocabulary",
    "Clinical_Sources",
    "Signal_Rules",
    "Signal_Rule_Groups",
    "Signal_Rule_Members",
    "Signal_Blockers",
}


def rows(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    raw = list(workbook[sheet_name].iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in raw[0]]
    return [
        dict(zip(headers, row))
        for row in raw[1:]
        if any(value not in (None, "") for value in row)
    ]


def norm(value: Any) -> str:
    return "" if value is None else str(value).strip()


def add_error(errors: list[str], condition: bool, message: str) -> None:
    if not condition:
        errors.append(message)


def unique_ids(errors: list[str], data: list[dict[str, Any]], field: str, label: str) -> set[str]:
    values = [norm(row.get(field)) for row in data]
    add_error(errors, all(values), f"{label}: blank {field}")
    duplicates = sorted(value for value, count in Counter(values).items() if value and count > 1)
    add_error(errors, not duplicates, f"{label}: duplicate {field}: {duplicates}")
    return set(values)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    data = {name: rows(workbook, name) for name in workbook.sheetnames if name in REQUIRED_SHEETS}
    errors: list[str] = []
    warnings: list[str] = []

    missing_sheets = sorted(REQUIRED_SHEETS - set(workbook.sheetnames))
    add_error(errors, not missing_sheets, f"Missing sheets: {missing_sheets}")
    if missing_sheets:
        print(json.dumps({"ready": False, "errors": errors}, indent=2))
        return 1

    signal_ids = unique_ids(errors, data["Signals"], "Signal_ID", "Signals")
    atom_ids = unique_ids(errors, data["Atoms"], "Atom_ID", "Atoms")
    bucket_ids = unique_ids(errors, data["Buckets"], "Reasoning_Bucket", "Buckets")
    combination_ids = unique_ids(errors, data["Combinations"], "Combination_ID", "Combinations")
    requirement_ids = unique_ids(
        errors, data["Combination_Requirements"], "Requirement_ID", "Combination_Requirements"
    )
    policy_ids = {norm(row["Priority_Policy_ID"]) for row in data["Priority_Policies"]}
    policy_keys = [
        (norm(row["Priority_Policy_ID"]), norm(row["Tier_Pattern"]))
        for row in data["Priority_Policies"]
    ]
    duplicate_policy_keys = sorted(key for key, count in Counter(policy_keys).items() if count > 1)
    add_error(errors, all(policy_id for policy_id, _ in policy_keys), "Priority_Policies: blank Priority_Policy_ID")
    add_error(errors, not duplicate_policy_keys, f"Priority_Policies: duplicate policy/tier pattern: {duplicate_policy_keys}")
    combination_rule_ids = unique_ids(errors, data["Combination_Rules"], "Combination_ID", "Combination_Rules")
    source_ids = unique_ids(errors, data["Clinical_Sources"], "Source_ID", "Clinical_Sources")

    for row in data["Signal_Atoms"]:
        add_error(errors, norm(row["Signal_ID"]) in signal_ids, f"Signal_Atoms unknown signal {row['Signal_ID']}")
        add_error(errors, norm(row["Atom_ID"]) in atom_ids, f"Signal_Atoms unknown atom {row['Atom_ID']}")
    for row in data["Terminology"]:
        add_error(errors, norm(row["Atom_ID"]) in atom_ids, f"Terminology unknown atom {row['Atom_ID']}")
        add_error(errors, bool(norm(row["Terminology_System"])), f"Terminology blank system for {row['Atom_ID']}")
        add_error(errors, bool(norm(row["Value"])), f"Terminology blank value for {row['Atom_ID']}")

    support_atoms = {norm(row["Atom_ID"]) for row in data["Signal_Atoms"]}
    blocker_atoms = {norm(row["Atom_ID"]) for row in data["Signal_Blockers"]}
    unconnected_atoms = sorted(atom_ids - support_atoms - blocker_atoms)
    add_error(errors, not unconnected_atoms, f"Atoms not connected to a signal or blocker: {unconnected_atoms}")

    mapped_by_signal: dict[str, set[str]] = defaultdict(set)
    for row in data["Signal_Atoms"]:
        mapped_by_signal[norm(row["Signal_ID"])].add(norm(row["Atom_ID"]))
    blockers_by_signal: dict[str, set[str]] = defaultdict(set)
    for row in data["Signal_Blockers"]:
        signal_id, atom_id = norm(row["Signal_ID"]), norm(row["Atom_ID"])
        add_error(errors, signal_id in signal_ids, f"Signal_Blockers unknown signal {signal_id}")
        add_error(errors, atom_id in atom_ids, f"Signal_Blockers unknown atom {atom_id}")
        blockers_by_signal[signal_id].add(atom_id)

    rule_ids = unique_ids(errors, data["Signal_Rules"], "Signal_ID", "Signal_Rules")
    add_error(errors, rule_ids == signal_ids, f"Signal_Rules coverage mismatch: missing={sorted(signal_ids-rule_ids)}, extra={sorted(rule_ids-signal_ids)}")
    group_keys: set[tuple[str, str]] = set()
    groups_by_signal: dict[str, set[str]] = defaultdict(set)
    for row in data["Signal_Rule_Groups"]:
        signal_id, group_id = norm(row["Signal_ID"]), norm(row["Group_ID"])
        key = (signal_id, group_id)
        add_error(errors, signal_id in signal_ids, f"Signal_Rule_Groups unknown signal {signal_id}")
        add_error(errors, bool(group_id), f"Signal_Rule_Groups blank group for {signal_id}")
        add_error(errors, key not in group_keys, f"Duplicate signal/group key {key}")
        group_keys.add(key)
        groups_by_signal[signal_id].add(group_id)
    for row in data["Signal_Rule_Groups"]:
        parent = norm(row["Parent_Group_ID"])
        if parent:
            add_error(errors, (norm(row["Signal_ID"]), parent) in group_keys, f"Unknown parent group {parent} for {row['Signal_ID']}")
    for row in data["Signal_Rules"]:
        add_error(errors, (norm(row["Signal_ID"]), norm(row["Root_Group_ID"])) in group_keys, f"Invalid root group for {row['Signal_ID']}")

    rule_atoms_by_signal: dict[str, set[str]] = defaultdict(set)
    for row in data["Signal_Rule_Members"]:
        signal_id, group_id = norm(row["Signal_ID"]), norm(row["Group_ID"])
        member_type, member_id = norm(row["Member_Type"]), norm(row["Member_ID"])
        add_error(errors, (signal_id, group_id) in group_keys, f"Signal_Rule_Members unknown group {signal_id}/{group_id}")
        if member_type == "ATOM":
            add_error(errors, member_id in atom_ids, f"Signal_Rule_Members unknown atom {member_id}")
            rule_atoms_by_signal[signal_id].add(member_id)
        elif member_type == "SIGNAL":
            add_error(errors, member_id in signal_ids, f"Signal_Rule_Members unknown signal member {member_id}")
        else:
            errors.append(f"Signal_Rule_Members invalid member type {member_type}")
    for signal_id in signal_ids:
        add_error(
            errors,
            mapped_by_signal[signal_id] == rule_atoms_by_signal[signal_id],
            f"Signal atom mismatch for {signal_id}: mapping={sorted(mapped_by_signal[signal_id])}, rules={sorted(rule_atoms_by_signal[signal_id])}",
        )

    for row in data["Combination_Requirements"]:
        add_error(errors, norm(row["Combination_ID"]) in combination_ids, f"Requirement unknown combination {row['Combination_ID']}")
    for row in data["Requirement_Buckets"]:
        add_error(errors, norm(row["Requirement_ID"]) in requirement_ids, f"Requirement_Buckets unknown requirement {row['Requirement_ID']}")
        add_error(errors, norm(row["Reasoning_Bucket"]) in bucket_ids, f"Requirement_Buckets unknown bucket {row['Reasoning_Bucket']}")
    for row in data["Requirement_Tiers"]:
        add_error(errors, norm(row["Requirement_ID"]) in requirement_ids, f"Requirement_Tiers unknown requirement {row['Requirement_ID']}")
    for row in data["Requirement_Signals"]:
        add_error(errors, norm(row["Requirement_ID"]) in requirement_ids, f"Requirement_Signals unknown requirement {row['Requirement_ID']}")
        add_error(errors, norm(row["Allowed_Signal_ID"]) in signal_ids, f"Requirement_Signals unknown signal {row['Allowed_Signal_ID']}")
    for row in data["Combinations"]:
        add_error(errors, norm(row["Priority_Policy_ID"]) in policy_ids, f"Combination unknown priority policy {row['Priority_Policy_ID']}")

    combination_group_keys: set[tuple[str, str]] = set()
    for row in data["Combination_Rule_Groups"]:
        combination_id, group_id = norm(row["Combination_ID"]), norm(row["Group_ID"])
        add_error(errors, combination_id in combination_rule_ids, f"Combination_Rule_Groups unknown rule {combination_id}")
        combination_group_keys.add((combination_id, group_id))
    for row in data["Combination_Rule_Members"]:
        combination_id, group_id = norm(row["Combination_ID"]), norm(row["Group_ID"])
        member_type, member_id = norm(row["Member_Type"]), norm(row["Member_ID"])
        add_error(errors, (combination_id, group_id) in combination_group_keys, f"Combination_Rule_Members unknown group {combination_id}/{group_id}")
        if member_type == "SIGNAL":
            add_error(errors, member_id in signal_ids, f"Combination_Rule_Members unknown signal {member_id}")
        elif member_type == "GROUP":
            add_error(errors, (combination_id, member_id) in combination_group_keys, f"Combination_Rule_Members unknown child group {combination_id}/{member_id}")
        else:
            errors.append(f"Combination_Rule_Members invalid member type {member_type}")
    for row in data["Combination_Rules"]:
        add_error(errors, norm(row["Combination_ID"]) in combination_ids, f"Combination rule unknown combination {row['Combination_ID']}")
        add_error(errors, norm(row["Priority_Policy_ID"]) in policy_ids, f"Combination rule unknown priority policy {row['Priority_Policy_ID']}")

    flat_combo_ids = {norm(row["Combination_ID"]) for row in data["Combination_Requirements"]}
    for combination_id in combination_ids:
        has_flat = combination_id in flat_combo_ids
        has_nested = combination_id in combination_rule_ids
        add_error(errors, has_flat != has_nested, f"Combination {combination_id} must have exactly one logic representation")

    cv: dict[str, set[str]] = defaultdict(set)
    for row in data["Controlled_Vocabulary"]:
        cv[norm(row["Vocabulary"])].add(norm(row["Allowed_Value"]))
    vocabulary_checks = {
        "Signals": {
            "Phenotype": "Phenotype",
            "Entity_Type": "Entity_Type",
            "Source_Type": "Source_Type",
            "Bucket_Class": "Bucket_Class",
            "Gate_Role": "Gate_Role",
            "Config_Action": "Config_Action",
            "Runtime_Executability": "Runtime_Executability",
        },
        "Signal_Atoms": {
            "Phenotype": "Phenotype",
            "Experiencer": "Experiencer",
            "Stage": "Stage",
            "Runtime_Executability": "Runtime_Executability",
        },
        "Combinations": {"Phenotype": "Phenotype", "Outcome": "Outcome"},
        "Combination_Requirements": {"Phenotype": "Phenotype", "Requirement_Kind": "Requirement_Kind"},
        "Priority_Policies": {"Priority_Class": "Priority_Class"},
        "Signal_Rule_Groups": {"Operator": "Rule_Group_Operator"},
        "Signal_Rule_Members": {"Member_Type": "Rule_Member_Type", "Experiencer": "Experiencer"},
        "Signal_Blockers": {"Block_Action": "Block_Action"},
        "Signal_Rules": {"Rule_Outcome": "Rule_Outcome", "Runtime_Executability": "Runtime_Executability"},
    }
    for sheet_name, checks in vocabulary_checks.items():
        for row in data[sheet_name]:
            for field, vocabulary in checks.items():
                value = norm(row.get(field))
                if value:
                    add_error(errors, value in cv[vocabulary], f"{sheet_name}.{field} value {value!r} outside {vocabulary}")

    duplicate_terminology = [
        key
        for key, count in Counter(
            (norm(row["Atom_ID"]), norm(row["Terminology_System"]), norm(row["Value"]).casefold())
            for row in data["Terminology"]
        ).items()
        if count > 1
    ]
    add_error(errors, not duplicate_terminology, f"Duplicate terminology keys: {duplicate_terminology[:20]}")

    structured_systems = {"ICD10", "ICD9", "CPT_HCPCS", "LOINC", "SNOMED_CT"}
    allowed_code_characters = {
        "ICD10": set("ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789.-"),
        "ICD9": set("ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789.-"),
        "CPT_HCPCS": set("ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789.-"),
        "LOINC": set("0123456789-"),
        "SNOMED_CT": set("0123456789"),
    }
    invalid_structured_rows: list[tuple[str, str, str]] = []
    for row in data["Terminology"]:
        system = norm(row["Terminology_System"]).upper()
        value = norm(row["Value"]).upper()
        value_class = norm(row["Value_Class"])
        if system in structured_systems:
            allowed = allowed_code_characters[system]
            valid_value = bool(value) and all(character in allowed for character in value) and any(character.isdigit() for character in value)
            if not valid_value or value_class != "STANDARD_CODE":
                invalid_structured_rows.append((norm(row["Atom_ID"]), system, norm(row["Value"])))
        elif system == "NLP" and value_class != "NLP_TERM":
            invalid_structured_rows.append((norm(row["Atom_ID"]), system, norm(row["Value"])))
    add_error(errors, not invalid_structured_rows, f"Invalid terminology system/value/class rows: {invalid_structured_rows[:20]}")

    terminology_by_atom: dict[str, list[dict[str, Any]]] = defaultdict(list)
    nlp_by_atom: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in data["Terminology"]:
        atom_id = norm(row["Atom_ID"])
        terminology_by_atom[atom_id].append(row)
        if norm(row["Terminology_System"]).upper() == "NLP":
            nlp_by_atom[atom_id].append(row)
    atoms_without_terms = sorted(atom_ids - terminology_by_atom.keys())
    atoms_without_nlp = sorted(atom_ids - nlp_by_atom.keys())
    add_error(errors, not atoms_without_terms, f"Atoms without terminology: {atoms_without_terms}")
    if atoms_without_nlp:
        warnings.append(f"Atoms without NLP terminology (may be structured/derived only): {atoms_without_nlp}")

    extraction_patterns = [row for row in data["Atoms"] if norm(row.get("Extraction_Pattern"))]
    add_error(errors, not extraction_patterns, "Atoms.Extraction_Pattern contains values; V4 prohibits regex extraction")

    clinical_source_refs: list[tuple[str, str]] = []
    for sheet_name in ("Signals", "Signal_Rules", "Signal_Blockers"):
        for row in data[sheet_name]:
            for source_id in norm(row.get("Clinical_Source_IDs")).replace(";", ",").split(","):
                if source_id.strip():
                    clinical_source_refs.append((sheet_name, source_id.strip()))
    unknown_source_refs = sorted({source_id for _, source_id in clinical_source_refs if source_id not in source_ids})
    add_error(errors, not unknown_source_refs, f"Unknown clinical source IDs: {unknown_source_refs}")

    report = {
        "ready": not errors,
        "workbook": str(workbook_path),
        "sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "row_counts": {name: len(sheet_rows) for name, sheet_rows in data.items()},
        "counts": {
            "signals": len(signal_ids),
            "atoms": len(atom_ids),
            "terminology": len(data["Terminology"]),
            "nlp_terminology": sum(len(values) for values in nlp_by_atom.values()),
            "atoms_with_nlp": len(nlp_by_atom),
            "combinations": len(combination_ids),
            "requirements": len(requirement_ids),
            "guardrails": len(data["Guardrails"]),
        },
        "errors": errors,
        "warnings": warnings,
    }
    print(json.dumps(report, indent=2, default=str))
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
