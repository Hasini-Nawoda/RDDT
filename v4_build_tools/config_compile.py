"""Compile the corrected normalized ATTRv workbook into deterministic V4 config.

This module is the only clinical-config ingestion path.  It does not inspect
or import any legacy configuration directory.  Workbook rows are retained as
JSON records with source sheet/row and workbook-hash provenance.
"""

from __future__ import annotations

import hashlib
import json
import math
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, MutableMapping, Optional, Sequence, Tuple

from .config_validation import (
    REQUIRED_SHEETS,
    ConfigValidationError,
    ValidationIssue,
    ValidationReport,
    validate_tables,
)

try:  # openpyxl is provided by the bundled workspace runtime.
    import openpyxl
except ImportError as exc:  # pragma: no cover - exercised only in a missing-runtime install
    openpyxl = None
    _OPENPYXL_IMPORT_ERROR = exc


COMPILER_VERSION = "v4-workbook-compiler-2026.09.19"
# The readiness-audited V3 workbook is the current corrected source.  A caller
# may still pass an explicit workbook path for a controlled recompile.
DEFAULT_WORKBOOK_FILENAME = "ATTRv_Normalized_Clinical_Filtering_Config_v3.xlsx"
DEFAULT_WORKBOOK_PATH = (
    Path(__file__).resolve().parent / "source" / DEFAULT_WORKBOOK_FILENAME
).resolve()

# Each required sheet has a strict header contract.  Extra columns are retained
# so a workbook can add audit metadata without this compiler dropping it.
REQUIRED_HEADERS: Dict[str, Tuple[str, ...]] = {
    "Signals": ("Phenotype", "Signal_ID", "Clinical_Feature", "Entity_Type", "Source_Type",
                 "Source_Specialty", "Source_Subdomain", "Reasoning_Bucket", "Bucket_Class", "Tier",
                 "Gate_Role", "Config_Action", "Canonical_Dedup_Group", "Algorithm_Role",
                 "Evidence_Stage", "Runtime_Executability", "Execution_Block_Reason", "Evidence_Mode",
                 "Atom_IDs", "Clinical_Logic", "Required_Qualifiers", "Guardrail_Notes", "Combination_IDs",
                 "Enabled", "Clinical_Source_IDs"),
    "Signal_Atoms": ("Phenotype", "Signal_ID", "Atom_ID", "Atom_Preferred_Name", "Experiencer", "Stage",
                     "Required_Qualifiers", "Signal_Logic", "Runtime_Executability", "Execution_Block_Reason",
                     "Evidence_Mode", "Can_Fire_From_This_Mapping"),
    "Atoms": ("Atom_ID", "Preferred_Clinical_Name", "Source_Type", "Source_Specialty", "Source_Subdomain",
              "Experiencer", "Stage", "Clinical_Meaning", "Context_Guard", "Contributing_Signals",
              "Extraction_Pattern"),
    "Terminology": ("Atom_ID", "Terminology_System", "Value", "Value_Class", "Review_Status",
                     "Can_Fire_Atom_Alone", "Context_Guard"),
    "Buckets": ("Phenotype", "Reasoning_Bucket", "Bucket_Class", "Counts_Independently", "Definition"),
    "Combinations": ("Phenotype", "Combination_ID", "Outcome", "Result_Route", "Priority_Policy_ID",
                      "Clinical_Rationale", "Validation_Class", "Enabled", "Origin", "Temporal_Policy", "Notes"),
    "Combination_Requirements": ("Phenotype", "Requirement_ID", "Combination_ID", "Requirement_Order",
                                  "Requirement_Kind", "Minimum_Count", "Context_Witness_Allowed",
                                  "Distinct_Lineage_Required", "Notes"),
    "Requirement_Buckets": ("Requirement_ID", "Reasoning_Bucket"),
    "Requirement_Tiers": ("Requirement_ID", "Allowed_Tier"),
    "Requirement_Signals": ("Requirement_ID", "Allowed_Signal_ID"),
    "Priority_Policies": ("Priority_Policy_ID", "Policy_Kind", "Tier_Pattern", "Priority_Class", "Notes"),
    "Combination_Rules": ("Phenotype", "Combination_ID", "Top_Operator", "Definition", "Priority_Policy_ID",
                           "Enabled"),
    "Combination_Rule_Groups": ("Combination_ID", "Group_ID", "Evaluation_Order", "Operator", "Minimum_Count",
                                 "Require_Independent_Lineage", "Notes"),
    "Combination_Rule_Members": ("Combination_ID", "Group_ID", "Member_Type", "Member_ID"),
    "Guardrails": ("Phenotype", "Guardrail_ID", "Guardrail_Action", "Route", "Does_Not_Negate_ATTRV",
                    "Clinical_Condition", "Message", "Enabled"),
    "Algorithm_Contract": ("Step", "Stage", "Action", "Invariant"),
    "Controlled_Vocabulary": ("Vocabulary", "Allowed_Value"),
    "Clinical_Sources": ("Source_ID", "Citation", "Year_or_Update", "URL", "Use_In_Config"),
}

OPTIONAL_HEADERS: Dict[str, Tuple[str, ...]] = {
    "Signal_Rules": ("Phenotype", "Signal_ID", "Root_Group_ID", "Rule_Outcome", "Three_Valued_Logic",
                      "Missing_Data_Policy", "Blocker_Policy", "Runtime_Executability", "Enabled",
                      "Clinical_Source_IDs", "Temporal_Policy"),
    "Signal_Rule_Groups": ("Phenotype", "Signal_ID", "Group_ID", "Parent_Group_ID", "Evaluation_Order",
                            "Operator", "Minimum_Count", "Linkage_Type", "Require_Independent_Lineage", "Notes"),
    "Signal_Rule_Members": ("Phenotype", "Signal_ID", "Group_ID", "Evaluation_Order", "Member_Type",
                             "Member_ID", "Required_Attributes", "Experiencer", "Member_Role"),
    "Signal_Blockers": ("Phenotype", "Signal_ID", "Blocker_ID", "Atom_ID", "Block_Action",
                         "Required_Attributes", "Evaluation_Semantics", "Clinical_Source_IDs"),
}

SHEET_TO_TABLE: Dict[str, str] = {
    "Signals": "signals",
    "Signal_Atoms": "signal_atoms",
    "Atoms": "atoms",
    "Terminology": "terminology",
    "Buckets": "buckets",
    "Combinations": "combinations",
    "Combination_Requirements": "combination_requirements",
    "Requirement_Buckets": "requirement_buckets",
    "Requirement_Tiers": "requirement_tiers",
    "Requirement_Signals": "requirement_signals",
    "Priority_Policies": "priority_policies",
    "Combination_Rules": "combination_rules",
    "Combination_Rule_Groups": "combination_rule_groups",
    "Combination_Rule_Members": "combination_rule_members",
    "Guardrails": "guardrails",
    "Algorithm_Contract": "algorithm_contract",
    "Controlled_Vocabulary": "controlled_vocabulary",
    "Clinical_Sources": "clinical_sources",
    "Signal_Rules": "signal_rules",
    "Signal_Rule_Groups": "signal_rule_groups",
    "Signal_Rule_Members": "signal_rule_members",
    "Signal_Blockers": "signal_blockers",
}


@dataclass
class CompiledBundle:
    output_dir: Path
    manifest: Dict[str, Any]
    tables: Dict[str, List[Dict[str, Any]]]
    validation_report: ValidationReport


def _json_safe(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        if value.is_integer():
            return int(value)
        return value
    if isinstance(value, Mapping):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    return value


def _key(header: Any) -> str:
    text = str(header or "").strip()
    pieces: List[str] = []
    for index, char in enumerate(text):
        if index and char.isupper() and pieces and pieces[-1] != "_":
            previous = text[index - 1]
            if previous.islower() or previous.isdigit():
                pieces.append("_")
            elif previous.isupper() and index >= 2 and text[index - 2].isupper():
                # Split before the new title-case word in e.g. IDValue,
                # while keeping two-letter suffixes such as IDs together.
                pieces.append("_")
        pieces.append(char.lower() if char.isalnum() else "_")
    collapsed: List[str] = []
    for char in pieces:
        if char == "_" and collapsed and collapsed[-1] == "_":
            continue
        collapsed.append(char)
    return "".join(collapsed).strip("_")


def _cell(value: Any) -> Any:
    value = _json_safe(value)
    if isinstance(value, str):
        return value.strip()
    return value


def _bool(value: Any, default: Optional[bool] = None) -> Optional[bool]:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    upper = str(value).strip().upper()
    if upper in {"TRUE", "YES", "Y", "1"}:
        return True
    if upper in {"FALSE", "NO", "N", "0"}:
        return False
    return default


def _int(value: Any, default: Optional[int] = None) -> Optional[int]:
    if value is None or value == "":
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _list_cell(value: Any) -> List[str]:
    if value is None or value == "":
        return []
    return [part.strip() for part in str(value).replace(",", ";").split(";") if part.strip()]


def _row_dict(headers: Sequence[str], values: Sequence[Any], sheet: str, row_number: int, workbook_hash: str) -> Dict[str, Any]:
    row = {header: _cell(values[index] if index < len(values) else None) for index, header in enumerate(headers)}
    row["_source_sheet"] = sheet
    row["_source_row"] = row_number
    row["_workbook_hash"] = workbook_hash
    return row


def _read_sheet(ws: Any, sheet: str, workbook_hash: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    values = list(ws.iter_rows(values_only=True))
    while values and not any(value is not None and str(value).strip() for value in values[0]):
        values.pop(0)
    if not values:
        return [], []
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    rows: List[Dict[str, Any]] = []
    for row_number, row_values in enumerate(values[1:], start=2):
        if not any(value is not None and str(value).strip() for value in row_values):
            continue
        rows.append(_row_dict(headers, row_values, sheet, row_number, workbook_hash))
    return headers, rows


def _snake_rows(rows: Sequence[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    result: List[Dict[str, Any]] = []
    for row in rows:
        converted: Dict[str, Any] = {}
        for field, value in row.items():
            converted[_key(field) if not field.startswith("_") else field] = _json_safe(value)
        result.append(converted)
    return result


def _sort_rows(table: str, rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    preferred: Dict[str, Tuple[str, ...]] = {
        "signals": ("phenotype", "signal_id"), "signal_atoms": ("signal_id", "atom_id"),
        "atoms": ("atom_id",), "terminology": ("atom_id", "terminology_system", "value"),
        "buckets": ("phenotype", "reasoning_bucket"), "combinations": ("phenotype", "combination_id"),
        "combination_requirements": ("combination_id", "requirement_order", "requirement_id"),
        "requirement_buckets": ("requirement_id", "reasoning_bucket"),
        "requirement_tiers": ("requirement_id", "allowed_tier"),
        "requirement_signals": ("requirement_id", "allowed_signal_id"),
        "priority_policies": ("priority_policy_id", "tier_pattern"),
        "combination_rules": ("phenotype", "combination_id"),
        "combination_rule_groups": ("combination_id", "evaluation_order", "group_id"),
        "combination_rule_members": ("combination_id", "group_id", "member_type", "member_id"),
        "guardrails": ("phenotype", "guardrail_id"), "clinical_sources": ("source_id",),
        "signal_rules": ("phenotype", "signal_id"),
        "signal_rule_groups": ("signal_id", "evaluation_order", "group_id"),
        "signal_rule_members": ("signal_id", "group_id", "evaluation_order", "member_id"),
        "signal_blockers": ("signal_id", "blocker_id"),
    }
    keys = preferred.get(table, tuple())
    def sort_key(row: Mapping[str, Any]) -> Tuple[str, ...]:
        return tuple(str(row.get(k) if row.get(k) is not None else "") for k in keys) + (str(row.get("_source_row", "")),)
    return sorted((dict(row) for row in rows), key=sort_key)


def _canonical_json(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode("utf-8")


def _table_file(table: str) -> str:
    return f"{table}.json"


def _add_issue(report: ValidationReport, issue: ValidationIssue) -> None:
    if issue.severity.upper() == "ERROR":
        report.errors.append(issue)
    else:
        report.warnings.append(issue)


def _attach_relationships(tables: MutableMapping[str, List[Dict[str, Any]]]) -> None:
    """Attach normalized child sheets to parent rows without replacing child tables."""
    atoms_by_id: Dict[str, Dict[str, Any]] = {str(row.get("atom_id")): row for row in tables.get("atoms", [])}
    term_count: Dict[str, int] = {}
    terms_by_atom: Dict[str, List[Dict[str, Any]]] = {}
    for row in tables.get("terminology", []):
        atom_id = str(row.get("atom_id"))
        term_count[atom_id] = term_count.get(atom_id, 0) + 1
        terms_by_atom.setdefault(atom_id, []).append(row)
    for row in tables.get("atoms", []):
        row["terminology_row_count"] = term_count.get(str(row.get("atom_id")), 0)
        row["terminology"] = _sort_rows("terminology", terms_by_atom.get(str(row.get("atom_id")), []))

    mappings_by_signal: Dict[str, List[Dict[str, Any]]] = {}
    for row in tables.get("signal_atoms", []):
        mappings_by_signal.setdefault(str(row.get("signal_id")), []).append(row)
    for row in tables.get("signals", []):
        row["mappings"] = mappings_by_signal.get(str(row.get("signal_id")), [])
        row["atom_ids_authoritative"] = [str(x.get("atom_id")) for x in row["mappings"]]

    req_by_id: Dict[str, Dict[str, Any]] = {str(row.get("requirement_id")): row for row in tables.get("combination_requirements", [])}
    for row in req_by_id.values():
        rid = str(row.get("requirement_id"))
        row["allowed_buckets"] = sorted({str(x.get("reasoning_bucket")) for x in tables.get("requirement_buckets", []) if str(x.get("requirement_id")) == rid})
        row["allowed_tiers"] = sorted({_int(x.get("allowed_tier")) for x in tables.get("requirement_tiers", []) if str(x.get("requirement_id")) == rid and _int(x.get("allowed_tier")) is not None})
        row["allowed_signal_ids"] = sorted({str(x.get("allowed_signal_id")) for x in tables.get("requirement_signals", []) if str(x.get("requirement_id")) == rid})
    req_by_combo: Dict[str, List[Dict[str, Any]]] = {}
    for row in req_by_id.values():
        req_by_combo.setdefault(str(row.get("combination_id")), []).append(row)
    for row in tables.get("combinations", []):
        row["requirements"] = sorted(req_by_combo.get(str(row.get("combination_id")), []), key=lambda x: (_int(x.get("requirement_order"), 0) or 0, str(x.get("requirement_id"))))

    members_by_group: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
    for row in tables.get("combination_rule_members", []):
        members_by_group.setdefault((str(row.get("combination_id")), str(row.get("group_id"))), []).append(row)
    for row in tables.get("combination_rule_groups", []):
        row["members"] = _sort_rows("combination_rule_members", members_by_group.get((str(row.get("combination_id")), str(row.get("group_id"))), []))
    groups_by_rule: Dict[str, List[Dict[str, Any]]] = {}
    for row in tables.get("combination_rule_groups", []):
        groups_by_rule.setdefault(str(row.get("combination_id")), []).append(row)
    for row in tables.get("combination_rules", []):
        row["groups"] = sorted(groups_by_rule.get(str(row.get("combination_id")), []), key=lambda x: (_int(x.get("evaluation_order"), 0) or 0, str(x.get("group_id"))))


def _load_raw_workbook(workbook_path: Path) -> Tuple[Dict[str, List[Dict[str, Any]]], Dict[str, int], str, Set[str], Dict[str, List[str]]]:
    if openpyxl is None:  # pragma: no cover
        raise RuntimeError("openpyxl is required to compile the workbook") from _OPENPYXL_IMPORT_ERROR
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    workbook_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    present = set(workbook.sheetnames)
    raw: Dict[str, List[Dict[str, Any]]] = {}
    counts: Dict[str, int] = {}
    headers_by_sheet: Dict[str, List[str]] = {}
    for sheet in list(REQUIRED_HEADERS) + list(OPTIONAL_HEADERS):
        if sheet not in present:
            continue
        headers, rows = _read_sheet(workbook[sheet], sheet, workbook_hash)
        raw[sheet] = rows
        counts[sheet] = len(rows)
        headers_by_sheet[sheet] = headers
    return raw, counts, workbook_hash, present, headers_by_sheet


def _validate_headers(headers_by_sheet: Mapping[str, Sequence[str]], present: Iterable[str], report: ValidationReport) -> None:
    present_set = set(present)
    for sheet, required in REQUIRED_HEADERS.items():
        if sheet not in present_set:
            continue
        actual = set(headers_by_sheet.get(sheet, []))
        missing = [header for header in required if header not in actual]
        if missing:
            report.add_error("MISSING_REQUIRED_HEADER", f"Missing required header(s): {', '.join(missing)}", sheet=sheet)
        if len(headers_by_sheet.get(sheet, [])) != len(set(headers_by_sheet.get(sheet, []))):
            report.add_error("DUPLICATE_HEADER", "Duplicate header names are not deterministic", sheet=sheet)
    for sheet, required in OPTIONAL_HEADERS.items():
        if sheet not in present_set:
            continue
        actual = set(headers_by_sheet.get(sheet, []))
        missing = [header for header in required if header not in actual]
        if missing:
            report.add_error("MISSING_REQUIRED_HEADER", f"Missing required header(s): {', '.join(missing)}", sheet=sheet)


def compile_workbook(
    workbook_path: Optional[str | Path] = None,
    output_dir: Optional[str | Path] = None,
    *,
    fail_on_error: bool = True,
    strict: Optional[bool] = None,
) -> CompiledBundle:
    """Compile a normalized workbook and emit deterministic JSON artifacts."""
    if strict is not None:
        fail_on_error = strict
    source = Path(workbook_path or DEFAULT_WORKBOOK_PATH).expanduser().resolve()
    # Build output stays outside the deployable runtime. A separate migration
    # step copies the validated logical tables into v4/config's categorized
    # layout for Snowflake consumption.
    out = Path(output_dir or Path(__file__).resolve().parent / "compiled_flat").expanduser().resolve()
    out.mkdir(parents=True, exist_ok=True)
    raw_by_sheet, row_counts, workbook_hash, present, headers = _load_raw_workbook(source)
    report = ValidationReport(workbook_sha256=workbook_hash)
    _validate_headers(headers, present, report)

    for sheet in REQUIRED_SHEETS:
        if sheet not in present:
            continue
        # validate_tables consumes original workbook headers, preserving clear
        # sheet/field names in its diagnostics.
    validation_tables = {sheet: rows for sheet, rows in raw_by_sheet.items()}
    structural = validate_tables(validation_tables, workbook_sha256=workbook_hash, sheets_present=present)
    report.errors.extend(structural.errors)
    report.warnings.extend(structural.warnings)

    tables: Dict[str, List[Dict[str, Any]]] = {}
    for sheet, rows in raw_by_sheet.items():
        table = SHEET_TO_TABLE.get(sheet)
        if table is not None:
            tables[table] = _sort_rows(table, _snake_rows(rows))
    _attach_relationships(tables)
    # Re-sort after attaching nested child rows so output never depends on
    # Python dictionary insertion order.
    for table in list(tables):
        tables[table] = _sort_rows(table, tables[table])

    phenotype_values = sorted({str(row.get("phenotype")) for row in tables.get("signals", []) if row.get("phenotype")})
    phenotype = phenotype_values[0] if len(phenotype_values) == 1 else (phenotype_values[0] if phenotype_values else "ATTRV")
    if len(phenotype_values) > 1:
        report.add_error("INCONSISTENT_PHENOTYPE", f"Multiple phenotype values found: {phenotype_values}", sheet="Signals")

    # Emit payload files first.  The compiled hash intentionally excludes the
    # manifest and validation report, which themselves contain that hash.
    payload_bytes: Dict[str, bytes] = {}
    for table in sorted(tables):
        payload_bytes[_table_file(table)] = _canonical_json(tables[table])
    compiled_hash_input = b"".join(name.encode("utf-8") + b"\0" + payload_bytes[name] for name in sorted(payload_bytes))
    compiled_hash = hashlib.sha256(compiled_hash_input).hexdigest()
    manifest: Dict[str, Any] = {
        "manifest_version": 1,
        "compiler_version": COMPILER_VERSION,
        "phenotype": phenotype,
        "workbook_filename": source.name,
        "workbook_sha256": workbook_hash,
        "source_workbook_sha256": workbook_hash,
        "compiled_config_sha256": compiled_hash,
        "compiled_config_hash": compiled_hash,
        "sheet_row_counts": {key: row_counts[key] for key in sorted(row_counts)},
        "required_sheets": list(REQUIRED_SHEETS),
        "optional_sheets_present": sorted(set(present) & set(OPTIONAL_HEADERS)),
        "tables": {table: _table_file(table) for table in sorted(tables)},
        "clinical_source_of_truth": "normalized workbook only; no legacy clinical fallback",
    }
    validation_report = report.to_dict()
    validation_report["compiler_version"] = COMPILER_VERSION
    validation_report["workbook_filename"] = source.name
    validation_report["compiled_config_sha256"] = compiled_hash

    expected_json = set(payload_bytes) | {"manifest.json", "validation_report.json"}
    for stale in out.glob("*.json"):
        if stale.name not in expected_json:
            stale.unlink()
    for name, data in payload_bytes.items():
        (out / name).write_bytes(data)
    (out / "manifest.json").write_bytes(_canonical_json(manifest))
    (out / "validation_report.json").write_bytes(_canonical_json(validation_report))

    bundle = CompiledBundle(output_dir=out, manifest=manifest, tables=tables, validation_report=report)
    if fail_on_error and not report.valid:
        first = report.errors[0]
        raise ConfigValidationError(
            f"Workbook configuration validation failed ({len(report.errors)} error(s)); "
            f"first: {first.code}: {first.message}", report
        )
    return bundle


def compile_to_bundle(*args: Any, **kwargs: Any) -> CompiledBundle:
    """Backward-compatible descriptive alias for :func:`compile_workbook`."""
    return compile_workbook(*args, **kwargs)


# Stable descriptive aliases used by orchestration/tests.
compile_config = compile_workbook
compile_workbook_config = compile_workbook
ConfigBundle = CompiledBundle
